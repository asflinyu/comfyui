# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
wb.remove(wb.active)
ws = wb.create_sheet("GEO推广文案汇总")

header_fill = PatternFill("solid", fgColor="70AD47")
header_font = Font(bold=True, color="FFFFFF", size=12, name="微软雅黑")
body_font = Font(size=11, name="微软雅黑")
title_font = Font(bold=True, size=14, name="微软雅黑")
industry_font = Font(bold=True, size=12, color="FFFFFF", name="微软雅黑")
industry_fill = PatternFill("solid", fgColor="70AD47")
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
left_wrap = Alignment(wrap_text=True, vertical="center", horizontal="left")

headers = ["#", "时间", "画面", "台词", "反转"]

# 总标题
ws.merge_cells("A1:E1")
c = ws.cell(row=1, column=1, value="GEO推广短视频文案汇总")
c.font = title_font
c.alignment = wrap
ws.row_dimensions[1].height = 32

# 表头
for col, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = wrap
    c.border = border

current_row = 3

data1 = [
    [1, "0-2s", "老王跪店门口磕头求顾客进店", "求您了，免费吃！", "顾客绕道走，嫌脏"],
    [2, "2-4s", "老王拎汽油桶要烧店转让", "不干了！", "朋友一脚踹门：\"烧啥！用这个\""],
    [3, "4-6s", "朋友亮GEO推广后台，订单999+", "让AI替你拉客", "老王冷笑：\"又是骗子\""],
    [4, "6-8s", "顾客问AI\"火锅推荐\"，AI推别家", "推荐张三火锅", "老王：\"看吧，没用\""],
    [5, "8-10s", "顾客补一句\"那家排队3小时，第二呢？\"AI推老王", "第二推荐老王", "老王愣住"],
    [6, "10-12s", "门口排长队，老王想关门歇业", "今天太累了", "顾客堵门不让关：\"等了2小时！\""],
    [7, "12-14s", "一年后，\"老王火锅·第8家分店\"挂牌", "媒体采访", "老王：\"我以前想烧店\""],
    [8, "14-15s", "老王对镜头比中指变比心", "现在别人求我开门", "全场反转"],
]

data2 = [
    [1, "0-2s", "老李给客户跪递方案，客户当场撕掉", "再考虑考虑", "老李捡方案手抖"],
    [2, "2-4s", "老李准备注销公司，盖公章", "认输", "妻子夺章：\"试试GEO\""],
    [3, "4-6s", "老李：\"AI能拉客户？扯淡\"", "不信", "妻子亮同行订单截图：\"人家8单/月\""],
    [4, "6-8s", "业主问AI\"装修推荐\"，AI推老李", "案例多", "老李：\"真的假的\""],
    [5, "8-10s", "电话响，老李以为是骚扰挂掉", "又是推销", "一看是业主，狂打回去"],
    [6, "10-12s", "日历排满到明年3月，老李想接低价单", "凑数", "妻子：\"别接，挑高端\""],
    [7, "12-14s", "老李把低价单全退，墙上挂\"只接5万+\"", "挑客户", "客户反而求他接"],
    [8, "14-15s", "老李翘腿喝茶对镜头", "以前求单，现在单求我", "身份彻底反转"],
]

data3 = [
    [1, "0-2s", "陈医生对空椅练手，护士打盹", "再不接单要关了", "门铃响，他冲过去"],
    [2, "2-4s", "来人不是看牙，是问奶茶店路", "请问奶茶店在哪", "陈医生指错方向赶人"],
    [3, "4-6s", "陈医生怒开通GEO推广", "让AI把我推到嘴边", "护士：\"万一没人来呢\""],
    [4, "6-8s", "患者问AI\"种牙推荐\"，AI推别家", "推荐XX口腔", "陈医生：\"果然没用\""],
    [5, "8-10s", "患者说\"那家太贵，第二呢？\"AI推陈医生", "陈医生性价比高", "陈医生跳起来"],
    [6, "10-12s", "预约排到3个月后，陈医生想休息", "今晚不加班了", "护士：\"又约了5台，跑不了\""],
    [7, "12-14s", "陈医生黑眼圈连做10台，累瘫", "要死了", "看到账单笑出声"],
    [8, "14-15s", "陈医生对镜头", "累，但月入翻了20倍", "苦→甜反转"],
]

sections = [
    ("文案1：火锅店老王（餐饮）", data1),
    ("文案2：装修公司老李（家装）", data2),
    ("文案3：牙科陈医生（医疗）", data3),
]

def write_section(ws, row, title, rows):
    # 行业标题行
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    c = ws.cell(row=row, column=1, value=title)
    c.fill = industry_fill
    c.font = industry_font
    c.alignment = left_wrap
    ws.row_dimensions[row].height = 26
    row += 1
    # 数据
    for r in rows:
        for col, v in enumerate(r, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = body_font
            c.border = border
            c.alignment = wrap if col in (1, 2) else left_wrap
        ws.row_dimensions[row].height = 32
        row += 1
    # 分隔空白行
    row += 1
    return row

for title, rows in sections:
    current_row = write_section(ws, current_row, title, rows)

widths = [6, 10, 42, 26, 40]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A3"

out = "/root/ComfyUI/GEO推广文案.xlsx"
wb.save(out)
print("saved:", out)
